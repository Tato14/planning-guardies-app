"""
Plantilles mensuals per setmanes completes.

Cada planning mensual va del 1r dilluns del mes al diumenge anterior al 1r
dilluns del mes següent (la setmana pertany al mes del seu dilluns).

Funcions principals
  build_input_template(master, any, mes, ...)  -> llibre d'ENTRADA del mes
      Copia les pestanyes Radiòlegs i Configuració del fitxer mestre i crea
      una pestanya Vacances nova amb una columna per data del període, els
      festius de Catalunya ja marcats i 2 dies de context del període anterior.
  build_output_template(base, any, mes, ...)   -> PLANNING buit del mes
      Una pestanya per setmana (dl→dg) amb les dates posades.
  festius_catalunya(any)                       -> festius a marcar

Ús des de línia d'ordres (genera un rang de mesos):
  python3 period_templates.py <entrada_mestra.xlsx> <plantilla_planning.xlsx> <carpeta> 2026-11 2027-12 \
      [--pont 2026-11-01] [--primer "Cognoms, Nom"]
"""
from __future__ import annotations

import datetime
import io
import os

import openpyxl
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from planning_generator import (
    DIES_CA, DIES_REINCORPORACIO, MESOS_CA, Meta, _es_actiu, _es_fila_llegenda,
    fmt_dia, month_label, next_month, period_for_month, prepare_planning_sheets,
)

# ============================================================
# Festius
# ============================================================
# Calendaris oficials verificats (Departament d'Empresa i Treball):
#   2026: Ordre EMT/66/2025
#   2027: Ordre EMT/52/2026 (DOGC núm. 9637, 1/4/2026)
FESTIUS_OFICIALS = {
    2026: [(1, 1), (1, 6), (4, 3), (4, 6), (5, 1), (6, 24), (8, 15), (9, 11),
           (10, 12), (12, 8), (12, 25), (12, 26)],
    2027: [(1, 1), (1, 6), (3, 26), (3, 29), (5, 1), (6, 24), (9, 11), (10, 12),
           (11, 1), (12, 6), (12, 8), (12, 25)],
}


def _pasqua(year: int) -> datetime.date:
    """Diumenge de Pasqua (algorisme gregorià anònim)."""
    a = year % 19
    b, c = divmod(year, 100)
    d, e = divmod(b, 4)
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = divmod(c, 4)
    l_ = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l_) // 451
    month = (h + l_ - 7 * m + 114) // 31
    day = ((h + l_ - 7 * m + 114) % 31) + 1
    return datetime.date(year, month, day)


def festius_tradicionals(year: int) -> list:
    """Els 14 festius habituals de Catalunya (inclosos els que cauen en diumenge)."""
    p = _pasqua(year)
    fixos = [(1, 1), (1, 6), (5, 1), (6, 24), (8, 15), (9, 11), (10, 12),
             (11, 1), (12, 6), (12, 8), (12, 25), (12, 26)]
    out = [datetime.date(year, m, d) for m, d in fixos]
    out += [p - datetime.timedelta(days=2), p + datetime.timedelta(days=1)]  # Div. Sant, Dl. Pasqua
    return sorted(out)


def festius_catalunya(year: int) -> tuple[list, bool]:
    """Festius a marcar a la plantilla i si l'any està verificat.

    Es marquen els 14 festius habituals: els oficials de l'any més els que cauen
    en diumenge (no afecten els torns, però s'hi veuen). Per als anys sense
    calendari publicat el resultat és provisional.
    """
    return festius_tradicionals(year), year in FESTIUS_OFICIALS


def festius_entre(start: datetime.date, end: datetime.date) -> tuple[list, bool]:
    out, verificat = [], True
    for y in range(start.year, end.year + 1):
        f, ok = festius_catalunya(y)
        out += [d for d in f if start <= d <= end]
        if any(start <= d <= end for d in f) and not ok:
            verificat = False
    return sorted(out), verificat


# ============================================================
# Estils
# ============================================================
BLAU = '4472C4'
F_TITOL = Font(bold=True, size=11)
F_VALOR = Font(bold=True, size=11, color='1F3864')
F_PETIT = Font(italic=True, size=9, color='808080')
F_HDR = Font(bold=True, size=9, color='FFFFFF')
F_DIA = Font(bold=True, size=9, color='1F3864')
F_DIA_CTX = Font(bold=True, size=9, color='808080')
FILL_HDR = PatternFill('solid', fgColor=BLAU)
FILL_CTX_HDR = PatternFill('solid', fgColor='A6A6A6')
FILL_CTX = PatternFill('solid', fgColor='EDEDED')
FILL_CAPSET = PatternFill('solid', fgColor='EEF3FA')
FILL_FESTIU = PatternFill('solid', fgColor='FFE699')
FILL_META = PatternFill('solid', fgColor='F2F2F2')
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center')
THIN = Side(style='thin', color='BFBFBF')
THICK = Side(style='medium', color='1F3864')

CODI_COLORS = {          # color de fons per codi (format condicional)
    'V': 'C6EFCE',       # vacances  — verd
    'B': 'F8CBAD',       # baixa     — vermell suau
    'C': 'BDD7EE',       # congrés   — blau
    'G': 'FFD966',       # guàrdia externa — groc fosc
    'X': 'D9D9D9',       # informatiu — gris
}


# ============================================================
# Helpers
# ============================================================

def _load(master):
    if isinstance(master, openpyxl.Workbook):
        return master
    if isinstance(master, (bytes, bytearray)):
        return openpyxl.load_workbook(io.BytesIO(master))
    return openpyxl.load_workbook(master)


def radiologists_from_master(wb) -> list:
    """[(nom, actiu)] en l'ordre de la pestanya Radiòlegs (sense files de llegenda)."""
    ws = wb['Radiòlegs']
    out, seen = [], set()
    for r in range(4, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if not isinstance(name, str) or not name.strip():
            continue
        name = name.strip()
        if _es_fila_llegenda(name) or name in seen:
            continue
        seen.add(name)
        out.append((name, _es_actiu(ws.cell(row=r, column=5).value)))
    return out


def _period_dates(year, month, start=None, end=None):
    p_start, p_end = period_for_month(year, month)
    start = start or p_start
    end = end or p_end
    dates = [start + datetime.timedelta(days=i) for i in range((end - start).days + 1)]
    return start, end, dates


# ============================================================
# Plantilla d'ENTRADA
# ============================================================

def _build_vacances(ws, year, month, start, end, festius, names, start_radiologist, spare_rows=12):
    n_ctx = max(0, DIES_REINCORPORACIO - 1)
    ctx = [start - datetime.timedelta(days=k) for k in range(n_ctx, 0, -1)]
    dates = [start + datetime.timedelta(days=i) for i in range((end - start).days + 1)]
    all_dates = ctx + dates
    c0 = 2
    last_col = c0 + len(all_dates) - 1
    L = get_column_letter
    HDR_ROW, FEST_ROW, FIRST = 4, 5, 6
    last_row = FIRST + len(names) + spare_rows - 1
    n_weeks = len({d - datetime.timedelta(days=d.weekday()) for d in dates})

    # --- Fila 1: mes i primer de la roda ---
    ws['A1'] = 'Mes del planning:'
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=9)
    ws['B1'] = month_label(year, month)
    ws.merge_cells(start_row=1, start_column=10, end_row=1, end_column=17)
    ws['J1'] = 'Primer radiòleg de la roda →'
    ws.merge_cells(start_row=1, start_column=18, end_row=1, end_column=max(32, last_col))
    ws['R1'] = start_radiologist or None
    # --- Fila 2: període ---
    ws['A2'] = 'Període planificat:'
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=9)
    ws['B2'] = start
    ws.merge_cells(start_row=2, start_column=10, end_row=2, end_column=17)
    ws['J2'] = 'fins al'
    ws.merge_cells(start_row=2, start_column=18, end_row=2, end_column=25)
    ws['R2'] = end
    ws.merge_cells(start_row=2, start_column=26, end_row=2, end_column=max(32, last_col))
    pont = ' · comença amb dia pont' if start.weekday() != 0 else ''
    ws['Z2'] = f"{n_weeks} setmanes (dl→dg){pont}"
    for ref in ('A1', 'A2', 'J1', 'J2'):
        ws[ref].font = F_TITOL
        ws[ref].alignment = LEFT
    for ref in ('B1', 'R1', 'B2', 'R2'):
        ws[ref].font = F_VALOR
        ws[ref].alignment = LEFT
        ws[ref].fill = FILL_META
    ws['B2'].number_format = 'dd/mm/yyyy'
    ws['R2'].number_format = 'dd/mm/yyyy'
    ws['Z2'].font = F_PETIT

    # --- Files 3-5: dia de la setmana, data, festiu ---
    ws['A3'] = f'(columnes grises = últims {n_ctx} dies del període anterior)' if n_ctx else None
    ws['A3'].font = F_PETIT
    ws['A4'] = 'Radiòleg'
    ws['A4'].font = Font(bold=True, color='FFFFFF')
    ws['A4'].fill = FILL_HDR
    ws['A5'] = 'Festiu (marca F) →'
    ws['A5'].font = Font(bold=True, color='9C5700')
    fset = set(festius)
    for i, d in enumerate(all_dates):
        col = c0 + i
        is_ctx = d < start
        c3 = ws.cell(row=3, column=col, value=DIES_CA[d.weekday()])
        c3.alignment = CENTER
        c3.font = F_DIA_CTX if is_ctx else F_DIA
        if not is_ctx and d.weekday() >= 5:
            c3.fill = FILL_CAPSET
        c4 = ws.cell(row=HDR_ROW, column=col, value=d)
        c4.number_format = 'dd/mm'
        c4.alignment = CENTER
        c4.font = F_HDR
        c4.fill = FILL_CTX_HDR if is_ctx else FILL_HDR
        c5 = ws.cell(row=FEST_ROW, column=col)
        c5.alignment = CENTER
        c5.font = Font(bold=True, color='9C5700')
        if is_ctx:
            c5.value = None
            c5.fill = FILL_CTX
        elif d in fset:
            c5.value = 'F'
        # cos
        for r in range(FIRST, last_row + 1):
            cell = ws.cell(row=r, column=col)
            cell.alignment = CENTER
            if is_ctx:
                cell.fill = FILL_CTX
            elif d.weekday() >= 5:
                cell.fill = FILL_CAPSET
        # separador de setmanes (dilluns) i d'inici de període
        if (d.weekday() == 0 or d == start) and i > 0:
            for r in range(3, last_row + 1):
                cell = ws.cell(row=r, column=col)
                cell.border = Border(left=THICK, top=cell.border.top, bottom=cell.border.bottom)

    # --- Noms ---
    for k, (name, _actiu) in enumerate(names):
        c = ws.cell(row=FIRST + k, column=1, value=name)
        c.alignment = LEFT
    for r in range(FIRST, last_row + 1):
        ws.cell(row=r, column=1).border = Border(bottom=THIN)

    rng_body = f"B{FIRST}:{L(last_col)}{last_row}"
    rng_dates = f"B3:{L(last_col)}{last_row}"

    # --- Format condicional (l'ordre fixa la prioritat: primer els codis) ---
    for code, color in CODI_COLORS.items():
        ws.conditional_formatting.add(
            rng_body, CellIsRule(operator='equal', formula=[f'"{code}"'],
                                 fill=PatternFill('solid', fgColor=color), font=Font(bold=True)))
    # Professionals amb Actiu = "No" a Radiòlegs -> fila en gris
    ws.conditional_formatting.add(
        f"A{FIRST}:{L(last_col)}{last_row}",
        FormulaRule(formula=[f"COUNTIFS('Radiòlegs'!$A$4:$A$400,$A{FIRST},'Radiòlegs'!$E$4:$E$400,\"No\")>0"],
                    font=Font(italic=True, color='A6A6A6'), fill=PatternFill('solid', fgColor='F2F2F2')))
    # Columna festiva (qualsevol valor a la fila 5) -> groc, també a la capçalera
    ws.conditional_formatting.add(
        rng_dates, FormulaRule(formula=[f'B${FEST_ROW}<>""'], fill=FILL_FESTIU,
                               font=Font(bold=True, color='9C5700')))

    # --- Validacions ---
    dv = DataValidation(type='list', formula1='"V,B,C,G,X"', allow_blank=True, showDropDown=False,
                        showErrorMessage=True, showInputMessage=True,
                        errorTitle='Codi no vàlid',
                        error='Valors permesos: V (vacances), B (baixa), C (congrés), G (guàrdia externa), X (ja assignada).',
                        promptTitle='Codi',
                        prompt='V=vacances, B=baixa, C=congrés, G=guàrdia externa, X=ja assignada.')
    ws.add_data_validation(dv)
    dv.add(rng_body)
    dvf = DataValidation(type='list', formula1='"F"', allow_blank=True, showDropDown=False,
                         showErrorMessage=True, errorTitle='Festiu', error='Escriu F per marcar el dia com a festiu, o deixa-ho buit.')
    ws.add_data_validation(dvf)
    dvf.add(f"{L(c0 + len(ctx))}{FEST_ROW}:{L(last_col)}{FEST_ROW}")
    dvr = DataValidation(type='list', formula1="'Radiòlegs'!$A$4:$A$400", allow_blank=True,
                         showDropDown=False, showErrorMessage=False, showInputMessage=True,
                         promptTitle='Primer de la roda',
                         prompt="Tria el radiòleg que obre la roda aquest mes (el següent a l'últim del mes anterior).")
    ws.add_data_validation(dvr)
    dvr.add('R1')

    # --- Mides i panells ---
    ws.column_dimensions['A'].width = 38
    for col in range(c0, last_col + 1):
        ws.column_dimensions[L(col)].width = 5.7
    ws.row_dimensions[4].height = 18
    ws.freeze_panes = f"B{FIRST}"
    ws.sheet_view.zoomScale = 90


INSTRUCCIONS_ENTRADA = [
    ("PLANTILLA D'ENTRADA — Planning de guàrdies de Telediagnòstic", True),
    ("{label} · període {ini} → {fi} ({setmanes} setmanes){pont}", True),
    ("", False),
    ("COM VA EL PERÍODE", True),
    ("• Cada planning va per SETMANES COMPLETES, de dilluns a diumenge.", False),
    ("• Una setmana pertany al mes en què cau el seu dilluns. Per això el planning d'un mes pot acabar", False),
    ("  els primers dies del mes següent, i cap planning inclou dies del mes anterior.", False),
    ("• Així cada setmana és sencera dins d'un sol fitxer, i els recordatoris setmanals no queden partits.", False),
    ("", False),
    ("PESTANYES", True),
    ("1. Configuració — qui cobreix cada torn fix i els perfils especials. Normalment no cal tocar-la.", False),
    ("2. Radiòlegs — llista mestra. Columna B = rol. Columna E = \"Actiu aquest mes\":", False),
    ("   posa \"No\" per treure algú del planning d'aquest mes sense esborrar-lo (buit = Sí).", False),
    ("3. Vacances — una columna per dia del període.", False),
    ("", False),
    ("COM OMPLIR LA PESTANYA VACANCES", True),
    ("• Fila 1: el mes ja ve posat. Tria el primer radiòleg de la roda (el desplegable té tota la llista).", False),
    ("• Fila 5 (Festiu): els festius oficials de Catalunya ja hi són marcats amb F.", False),
    ("  Afegeix una F si cal tractar algun altre dia com a festiu, o esborra-la si no.", False),
    ("• A la fila de cada radiòleg, marca els dies que no pot fer guàrdia:", False),
    ("     V = Vacances      → bloqueja el dia + 2 dies de marge de reincorporació", False),
    ("     B = Baixa         → igual que V", False),
    ("     C = Congrés       → bloqueja només el dia", False),
    ("     G = Guàrdia en una altra institució → bloqueja el dia, l'anterior i el posterior", False),
    ("     X = Guàrdia ja assignada (només informatiu)", False),
    ("• Les columnes GRISES del principi són els últims dies del període anterior. Marca-hi V o B", False),
    ("  NOMÉS si l'absència ja venia d'abans: així el marge de reincorporació s'aplica també a l'inici", False),
    ("  d'aquest període. Aquests dies no es planifiquen.", False),
    ("• Si entra algú nou, afegeix-lo a Radiòlegs i escriu el seu nom en una fila buida del final de Vacances.", False),
    ("", False),
    ("DESPRÉS", True),
    ("• Puja aquest fitxer a l'app i clica \"Generar planning\".", False),
    ("• En acabar, l'app et deixa descarregar l'entrada del mes següent ja preparada:", False),
    ("  mateixos radiòlegs i Actiu, primer de la roda calculat i festius marcats.", False),
]


def _write_input_instructions(wb, label, start, end, n_weeks, verificat):
    if 'Instruccions' in wb.sheetnames:
        ws = wb['Instruccions']
        for row in ws.iter_rows():
            for c in row:
                c.value = None
    else:
        ws = wb.create_sheet('Instruccions', 0)
    pont = ' · comença amb dia pont' if start.weekday() != 0 else ''
    lines = list(INSTRUCCIONS_ENTRADA)
    if not verificat:
        lines += [("", False),
                  ("ATENCIÓ: el calendari oficial de festius d'algun any d'aquest període encara no estava "
                   "publicat; els festius marcats són provisionals. Reviseu-los.", True)]
    for i, (txt, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=txt.format(label=label, ini=fmt_dia(start, True),
                                                       fi=fmt_dia(end, True), setmanes=n_weeks, pont=pont))
        c.font = Font(bold=bold, size=12 if i == 1 else 11)
    ws.column_dimensions['A'].width = 110


def build_input_template(master, year, month, *, start=None, end=None,
                         start_radiologist='', festius=None, spare_rows=12):
    """Llibre d'entrada del mes a partir d'un fitxer mestre (path, bytes o Workbook)."""
    wb = _load(master)
    for req in ('Radiòlegs', 'Configuració'):
        if req not in wb.sheetnames:
            raise ValueError(f"Al fitxer mestre hi falta la pestanya «{req}».")
    start, end, dates = _period_dates(year, month, start, end)
    if festius is None:
        festius, verificat = festius_entre(start, end)
    else:
        verificat = True
    names = radiologists_from_master(wb)

    pos = wb.sheetnames.index('Vacances') if 'Vacances' in wb.sheetnames else len(wb.sheetnames)
    if 'Vacances' in wb.sheetnames:
        wb.remove(wb['Vacances'])
    ws = wb.create_sheet('Vacances', pos)
    _build_vacances(ws, year, month, start, end, festius, names, start_radiologist, spare_rows=spare_rows)
    n_weeks = len({d - datetime.timedelta(days=d.weekday()) for d in dates})
    _write_input_instructions(wb, month_label(year, month), start, end, n_weeks, verificat)
    wb.active = wb.sheetnames.index('Vacances')
    for s in wb.worksheets:
        s.sheet_view.tabSelected = (s.title == 'Vacances')
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    return wb


# ============================================================
# Plantilla del PLANNING (sortida)
# ============================================================

INSTRUCCIONS_PLANNING = [
    ("PLANNING DE GUÀRDIES — Telediagnòstic IDI", True),
    ("{label} · {ini} → {fi} ({setmanes} setmanes)", True),
    ("", False),
    ("Estructura", True),
    ("• Una pestanya per setmana COMPLETA, de dilluns a diumenge (nom: Setmana NN (dd mmm-dd mmm)).", False),
    ("• La setmana pertany al mes del seu dilluns: el planning pot acabar els primers dies del mes següent.", False),
    ("• Cada pestanya té 7 columnes (dl→dg) i 3 franges: 8-16, 16-20 i 20-08. La data és a la fila 3.", False),
    ("• Els dies fora del període (si n'hi ha) surten en gris i no s'han d'omplir.", False),
    ("", False),
    ("Convencions per cel·la", True),
    ("• \"N: Nom\" — torn de Neuro fix · \"B: Nom\" — torn de Body fix", False),
    ("• \"N i B: Nom\" — un sol radiòleg cobreix Neuro i Body", False),
    ("• \"Grup 1: Nom\" / \"Grup 2: Nom\" — torn repartit en dos grups", False),
    ("• Festius: capçalera groga; prefix \"FESTIU — Torn dia/nit\".", False),
    ("• [SUBSTITUIR — …] — torn fix que cal cobrir a mà (vacances, baixa, NO actiu…).", False),
    ("", False),
    ("No canvieu els noms de les pestanyes ni les files de les franges: els fluxos de Microsoft 365", False),
    ("llegeixen el fitxer automàticament (busquen les pestanyes que contenen \"Setmana\").", False),
]


def build_output_template(base, year, month, *, start=None, end=None, festius=None):
    """Planning buit del mes, amb una pestanya per setmana i les dates posades."""
    wb = _load(base)
    start, end, dates = _period_dates(year, month, start, end)
    if festius is None:
        festius, _ = festius_entre(start, end)
    fest_idx = [(d - start).days + 1 for d in festius if start <= d <= end]
    meta = Meta(year=year, month=month, festius=fest_idx, dates=dates)
    prepare_planning_sheets(wb, meta)
    if 'Instruccions' in wb.sheetnames:
        ws = wb['Instruccions']
        for row in ws.iter_rows():
            for c in row:
                c.value = None
        for i, (txt, bold) in enumerate(INSTRUCCIONS_PLANNING, start=1):
            c = ws.cell(row=i, column=1, value=txt.format(label=meta.label, ini=fmt_dia(start, True),
                                                           fi=fmt_dia(end, True), setmanes=meta.n_weeks))
            c.font = Font(bold=bold, size=12 if i == 1 else 11)
    return wb


# ============================================================
# Generació per lots (CLI)
# ============================================================

def month_range(first: tuple, last: tuple) -> list:
    out, (y, m) = [], first
    while (y, m) <= last:
        out.append((y, m))
        y, m = next_month(y, m)
    return out


def generate_batch(master_path, base_path, out_dir, first, last, pont=None, primer=''):
    os.makedirs(out_dir, exist_ok=True)
    fets = []
    for (y, m) in month_range(first, last):
        start = None
        if pont and period_for_month(y, m)[0] > pont and (pont.year, pont.month) == (y, m):
            start = pont
        nom = f"{y}-{m:02d} {MESOS_CA[m - 1].capitalize()}"
        wb_in = build_input_template(master_path, y, m, start=start,
                                     start_radiologist=primer if not fets else '')
        p_in = os.path.join(out_dir, f"{nom} - Entrada.xlsx")
        wb_in.save(p_in)
        wb_out = build_output_template(base_path, y, m, start=start)
        p_out = os.path.join(out_dir, f"{nom} - Planning buit.xlsx")
        wb_out.save(p_out)
        fets.append((p_in, p_out))
    return fets


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('mestre')
    ap.add_argument('plantilla_planning')
    ap.add_argument('carpeta')
    ap.add_argument('des_de', help='AAAA-MM')
    ap.add_argument('fins_a', help='AAAA-MM')
    ap.add_argument('--pont', help='Data d\'inici del primer període si cal un dia pont (AAAA-MM-DD)')
    ap.add_argument('--primer', default='', help='Primer radiòleg de la roda del primer mes')
    a = ap.parse_args()
    f = tuple(int(x) for x in a.des_de.split('-'))
    l_ = tuple(int(x) for x in a.fins_a.split('-'))
    pont = datetime.date.fromisoformat(a.pont) if a.pont else None
    for p_in, p_out in generate_batch(a.mestre, a.plantilla_planning, a.carpeta, f, l_, pont, a.primer):
        print('·', os.path.basename(p_in), '|', os.path.basename(p_out))
